"""The run's Excel workbook, laid out like the earlier verifier's report:
Record Summary, Other Locations, New Records Added, plus Field Details."""

import math
import re

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .fields import LABELS, SOCIAL_KEYS
from .normalize import format_phone

SUMMARY_COLUMNS = [
    "redcap_record_id", "facility_name", "updated_redcap_record", "manual_review_required",
    "found_multiple_locations", "new_records_added", "ai_assisted", "notes",
    "website_changed", "website_old", "website_new", "website_suggested",
    "phone_changed", "phone_old", "phone_new",
    "fax_changed", "fax_old", "fax_new",
    "email_changed", "email_old", "email_new",
    "social_changed", "social_old", "social_new",
    "address_changed", "address_old", "address_new",
    "name_changed", "name_old", "name_new",
]
OTHER_LOCATIONS_COLUMNS = ["record_id", "name", "street", "city", "state", "zip", "phone", "fax", "email",
                           "url", "found_by", "reason"]
CREATED_COLUMNS = ["record_id", "new_record_id", "name", "street", "city", "state", "zip", "county",
                   "phone", "fax", "email", "url", "found_by", "note"]
DETAIL_COLUMNS = ["record_id", "field", "in_redcap", "on_website", "result", "action", "found_by_ai",
                  "note", "all_values_found", "source_page"]

CATEGORIES = {
    "website": ("website",), "phone": ("phone",), "fax": ("fax",), "email": ("email",),
    "social": SOCIAL_KEYS, "address": ("street", "city", "state", "zip", "county"),
}


def _yn(value):
    return "Yes" if value else "No"


def _old_new(changes, record):
    if not changes:
        return "", ""
    if len(changes) == 1:
        return changes[0].on_file or "(NONE)", changes[0].new_value
    old = "; ".join(f"{LABELS[c.key]}: {c.on_file or '(NONE)'}" for c in changes)
    new = "; ".join(f"{LABELS[c.key]}: {c.new_value}" for c in changes)
    return old, new


def _change_line(f):
    if f.on_file:
        return f"The {LABELS[f.key]} was updated from {f.on_file} to {f.new_value}."
    return f"The {LABELS[f.key]} was added: {f.new_value}."


def _review_line(f):
    line = f"- {LABELS[f.key]}: {f.note or 'needs a decision'}"
    shown = [f"REDCap has: {f.on_file}"] if f.on_file else []
    if f.found:
        shown.append(f"website has: {', '.join(f.found[:6])}")
    return f"{line} ({'; '.join(shown)})" if shown else line


def _held_line(f):
    """A change found on the website but not written because the record is
    waiting for a person."""
    if f.on_file:
        return f"- {LABELS[f.key]}: change from {f.on_file} to {f.new_value}"
    return f"- {LABELS[f.key]}: add {f.new_value}"


def _fetch_problem(error):
    """A short, plain reason for a website that couldn't be opened."""
    if re.search(r"NameResolution|getaddrinfo|Name or service not known", error):
        return "the domain no longer exists"
    status = re.search(r"\b(4\d\d|5\d\d) (?:Client|Server) Error", error)
    if status:
        return {"404": "page not found (404)", "410": "page removed (410)",
                "403": "access denied (403)"}.get(status.group(1), f"server error {status.group(1)}")
    if "Timeout" in error or "timed out" in error:
        return "the site didn't respond in time"
    if "SSL" in error:
        return "the site's security certificate is invalid"
    return error.split(":")[0] or "unknown error"


def _notes(r, outcome, dry_run, suggest_websites=False):
    written, message = outcome
    if r.status in ("no website", "fetch failed"):
        if r.status == "no website":
            note = "Manual Review Required: no website on file for this record."
        else:
            note = f"Manual Review Required: the website on file couldn't be opened ({_fetch_problem(r.note)})."
        if r.suggestion:
            s = r.suggestion
            note += (f"\nSuggested website (please confirm; not written to REDCap): {s.url} "
                     f"({s.evidence}; {s.confidence} confidence)")
        elif suggest_websites:
            note += "\nWeb search found no website that matches this record's phone or address."
        return note

    sections = []
    reviews = [f for f in r.reviews if f.key != "name"]
    if r.blocked:
        sections.append("Manual Review Required: nothing was written for this record. Please decide:\n"
                        + "\n".join(_review_line(f) for f in reviews))
        if r.changes:
            sections.append("Also found on the website, waiting for the review above (not written yet):\n"
                            + "\n".join(_held_line(f) for f in r.changes))
    elif r.changes:
        label = "Updated:" if written else "Potential Update:"
        sections.append(label + "\n" + "\n".join(_change_line(f) for f in r.changes))
        if not dry_run and not written and message:
            sections.append(f"Manual Review Required: REDCap was not updated: {message}")
    name = next((f for f in r.fields if f.key == "name"), None)
    if name and name.status == "REVIEW":
        sections.append(f"Manual Review Required: Facility name (possible name change on the website, never "
                        f"auto-written; website shows {name.new_value!r})")
    if r.note and r.note not in "\n".join(sections):
        sections.append(f"Note: {r.note}")
    return "\n\n".join(sections) if any(s.startswith(("Manual", "Updated", "Potential")) for s in sections) \
        else "\n\n".join(["Record up to date: No Change Required"] + sections)


def summary_row(r, outcome, dry_run, suggest_websites=False):
    rec = r.record
    written, _ = outcome
    row = {
        "redcap_record_id": r.record_id,
        "facility_name": rec["name"],
        "updated_redcap_record": _yn(written),
        "manual_review_required": _yn(r.status in ("no website", "fetch failed") or bool(r.reviews)
                                      or (not dry_run and r.changes and not written)),
        "found_multiple_locations": _yn(bool(r.other_locations)),
        "new_records_added": ", ".join(o.new_record_id for o in r.other_locations
                                       if o.new_record_id and not o.new_record_id.startswith("(")),
        "ai_assisted": _yn(bool(r.ai_note)),
        "website_suggested": r.suggestion.url if r.suggestion else "",
        "website_old": rec["website"], "phone_old": rec["phone"], "fax_old": rec["fax"],
        "email_old": rec["email"], "name_old": rec["name"],
        "address_old": f"{rec['street']}, {rec['city']}, {rec['state']} {rec['zip']}".strip(", "),
        "notes": _notes(r, outcome, dry_run, suggest_websites),
    }
    changes = [] if r.blocked else r.changes
    for category, keys in CATEGORIES.items():
        mine = [c for c in changes if c.key in keys]
        old, new = _old_new(mine, rec)
        row[f"{category}_changed"] = _yn(bool(mine))
        row[f"{category}_new"] = new
        if category == "social":
            row["social_old"] = old
    for f in r.fields:
        if f.key == "website" and f.status == "REVIEW":
            row.update(website_changed="Yes", website_new=f.new_value)
        if f.key == "name" and f.status == "REVIEW":
            row.update(name_changed="Yes", name_new=f.new_value)
    row.setdefault("name_changed", "No")
    return row


def _action(f, r, written, dry_run):
    if not f.writes:
        return {"MATCH": "none (matches)", "NOT_FOUND": "none (not on website)"}.get(f.status, "manual review")
    if r.blocked:
        return "manual review (record has unclear fields)"
    if dry_run:
        return "would add" if f.status == "NEW" else "would update"
    return ("added" if f.status == "NEW" else "updated") if written else "not written"


def _sheet(wb, title, columns, rows, wrap=()):
    ws = wb.create_sheet(title)
    ws.append(columns)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in rows:
        ws.append([row.get(c, "") for c in columns])
    for line in ws.iter_rows(min_row=2):
        for cell in line:
            cell.alignment = Alignment(vertical="top")    # rows read across even when notes are tall
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions
    for i, col in enumerate(columns, start=1):
        letter = get_column_letter(i)
        if col in wrap:
            ws.column_dimensions[letter].width = 80
            for cell in ws[letter][1:]:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                # count wrapped lines too (about 85 characters fit in the
                # 80-wide column), so the end of a long note isn't cut off
                lines = sum(max(1, math.ceil(len(part) / 85)) for part in str(cell.value or "").split("\n"))
                ws.row_dimensions[cell.row].height = max(15, min(15 * lines, 409))
        else:
            width = max([len(col)] + [len(str(r.get(col, "")).split("\n")[0]) for r in rows[:500]])
            ws.column_dimensions[letter].width = min(max(width + 2, 10), 50)
        for cell in ws[letter][1:]:
            if cell.value in ("Yes", "No"):
                fill, font = _YES_NO_COLORS[(cell.value, col in _ATTENTION_COLUMNS)]
                cell.fill, cell.font = fill, font
                cell.alignment = Alignment(horizontal="center", vertical="top")
    return ws


# Yes/No cells: green for Yes, amber for a Yes that needs a person (manual
# review), grey for No. Same colors as Excel's Good / Neutral cell styles.
_ATTENTION_COLUMNS = {"manual_review_required"}
_YES_NO_COLORS = {
    ("Yes", False): (PatternFill("solid", fgColor="C6EFCE"), Font(color="006100", bold=True)),
    ("Yes", True): (PatternFill("solid", fgColor="FFEB9C"), Font(color="9C5700", bold=True)),
    ("No", False): (PatternFill("solid", fgColor="F2F2F2"), Font(color="808080")),
    ("No", True): (PatternFill("solid", fgColor="F2F2F2"), Font(color="808080")),
}


def _location_row(record_id, o):
    loc = o.location
    return {
        "record_id": record_id, "name": loc.name, "street": loc.street, "city": loc.city, "state": loc.state,
        "zip": loc.zip, "phone": format_phone(loc.phone) if loc.phone else "",
        "fax": format_phone(loc.fax) if loc.fax else "", "email": loc.email, "url": loc.url,
        "found_by": "AI" if loc.source == "AI" else "code", "reason": o.reason, "new_record_id": o.new_record_id,
        "county": o.county, "note": o.note,
    }


def write_report(path, results, outcomes, dry_run, suggest_websites=False):
    summary, others, created, details = [], [], [], []
    for r in results:
        outcome = outcomes.get(r.record_id, (False, ""))
        summary.append(summary_row(r, outcome, dry_run, suggest_websites))
        for o in r.other_locations:
            (created if o.new_record_id else others).append(_location_row(r.record_id, o))
        for f in r.fields:
            details.append({
                "record_id": r.record_id, "field": LABELS[f.key], "in_redcap": f.on_file,
                "on_website": f.new_value or "; ".join(f.found[:5]), "result": f.status,
                "action": _action(f, r, outcome[0], dry_run), "found_by_ai": _yn(f.ai), "note": f.note,
                "all_values_found": "\n".join(f.found), "source_page": f.source,
            })

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    _sheet(wb, "Record Summary", SUMMARY_COLUMNS, summary, wrap=("notes",))
    _sheet(wb, "Other Locations", OTHER_LOCATIONS_COLUMNS, others)
    _sheet(wb, "New Records Added", CREATED_COLUMNS, created)
    _sheet(wb, "Field Details", DETAIL_COLUMNS, details)
    wb.save(path)
