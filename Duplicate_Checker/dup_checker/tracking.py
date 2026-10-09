"""Fill the "Tracking Duplicates" tab of the duplicates workbook, in the same layout as manual reviews.

One row per record, in group order:
  - record kept (most complete):  its duplicates under "potential duplicate with?" and "which is true
                                  duplicate?", and its own ID under "most complete record"
  - duplicate records:            no row of their own - they're listed on the kept record's row
  - not a duplicate:              "Not a duplicate"
  - undecided:                    "Manual review required - possible duplicate with ..."
Name/address/other-changes columns are left blank for the reviewer. "validated date?" gets the run date
for records the app decided, and stays blank for "Manual review required" rows.
"notes (add to redcap)" holds exactly the comment --apply adds to REDCap; the app's explanation of each
decision goes in an extra "notes from automation" column.
"""
from __future__ import annotations

import io
import posixpath
from copy import copy
import re
import shutil
import zipfile
from datetime import date
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .checker import RecordResult
from .redcap import PlannedUpdate, describe_changes, plan_updates

SHEET = "Tracking Duplicates"
HEADERS = [
    "Redcap ID",
    "potential duplicate with? \n(tip: search for potential duplicates by name and address)",
    "which is true duplicate?",
    "which record is most complete record?\n (mark as not duplicate, verify info, ud",
    "records that are not actual duplicates",
    "Old name (previously in redcap)",
    "New name",
    "Old address (previously in redcap)",
    "New address",
    "Mailing Address",
    "other changes",
    "validated date?",
    "notes (add to redcap)",
    "notes from automation",
    "status",
]
NOTES_COL, AUTOMATION_COL, STATUS_COL = 13, 14, 15   # M: comment --apply puts in REDCap, N: app's explanation, O: status
STATUS_COLORS = {                                    # same colors as the REDCap change log
    "updated": "C6EFCE", "partly updated": "C6EFCE", "skipped": "FFEB9C", "failed": "F8CBAD",
    "need review": "F4B183",
}
SOURCE = {"Rule": "rule check", "AI": "AI review", "Website rule": "website check"}


def _decided_by(decided: str) -> str:
    return ", ".join(SOURCE.get(d, d) for d in decided.split("/") if d and d != "-") or "the app"


def _short(text: str, limit: int = 700) -> str:
    text = re.sub(r"\s+", " ", text or "").strip()
    return text if len(text) <= limit else text[:limit].rsplit(" ", 1)[0] + " ..."


def build_rows(results: list[RecordResult], run_date: date | None = None,
               redcap_outcome: list[PlannedUpdate] | None = None) -> list[list]:
    """Rows for the tab, in the order of `results` (group order).
    "notes (add to redcap)" is exactly the comment --apply adds to REDCap's Additional Comments (blank for
    records --apply doesn't change); "notes from automation" explains how the app decided.
    `redcap_outcome` (the plans after an --apply run) adds what actually happened in REDCap: records that
    were not updated are explained in "notes (add to redcap)", records that were updated are listed in
    "notes from automation"."""
    run_date = run_date or date.today()
    redcap_comment = {p.record_id: p.comment_to_add for p in plan_updates(results, run_date)}
    outcome_by_set: dict[str, list[PlannedUpdate]] = {}
    for p in redcap_outcome or []:
        outcome_by_set.setdefault(p.primary_id, []).append(p)
    rows = []
    for res in results:
        r, rec = res.record, res.recommendation
        gid = r.group_id
        if rec.startswith("Duplicate of"):
            continue                                   # listed on the kept record's row
        dup_with = ", ".join(res.duplicate_with)
        if rec.startswith("Duplicate - keep"):
            cols = [dup_with, dup_with, r.record_id]
            note = (f"Group {gid}: true duplicate. Keep this record (most complete); duplicate record(s): "
                    f"{dup_with}. Decided by {_decided_by(res.decided_by)}: {_short(res.reason)}")
        elif rec == "Not a duplicate":
            cols = ["Not a duplicate", "", ""]
            note = f"Group {gid}: not a duplicate. Decided by {_decided_by(res.decided_by)}: {_short(res.reason)}"
        else:
            cols = [f"Manual review required - possible duplicate with {dup_with}".strip(" -"), "", ""]
            note = (f"Group {gid}: MANUAL REVIEW REQUIRED - the app could not decide whether this record is a "
                    f"duplicate of {dup_with or 'the other records'}. {_short(res.reason)}")
        # "validated date?": the run date for decided records; blank while a person still has to decide
        validated = "" if rec == "Needs review" else run_date
        redcap_note = redcap_comment.get(r.record_id, "")
        status = "need review" if rec == "Needs review" else ""
        if r.record_id in outcome_by_set:
            redcap_note, note = _with_redcap_outcome(outcome_by_set[r.record_id], r.record_id, redcap_note,
                                                     note, run_date)
            status = _set_status(outcome_by_set[r.record_id])
        rows.append([_number(r.record_id), *map(_number, cols), "", "", "", "", "", "", "", validated,
                     redcap_note, note, status])
    return rows


def _set_status(plans: list[PlannedUpdate]) -> str:
    """One status for a kept record and its duplicates."""
    statuses = {p.status for p in plans}
    if statuses & {"failed", "sending"}:
        return "failed"
    if statuses == {"updated"}:
        return "updated"
    if "updated" in statuses:
        return "partly updated"
    return "skipped" if statuses == {"skipped"} else ""


def _with_redcap_outcome(plans: list[PlannedUpdate], kept_id: str, redcap_note: str, note: str,
                         run_date: date) -> tuple[str, str]:
    """Add what --apply did to a kept record's row (the row covers its duplicates too)."""
    plans = sorted(plans, key=lambda p: p.role != "kept")          # the kept record first
    kept = next((p for p in plans if p.record_id == kept_id), None)
    # the comment only went into REDCap if the kept record was actually updated
    lines = [redcap_note] if kept is not None and kept.status == "updated" else []
    lines += [f"Record {p.record_id} not updated in REDCap - {p.note}" for p in plans if p.status == "skipped"]
    updated = [f"record {p.record_id} ({describe_changes(p)})" for p in plans if p.status == "updated"]
    failed = [f"record {p.record_id} ({p.note})" for p in plans if p.status in ("failed", "sending")]
    if updated:
        note += f" REDCap updated on {run_date:%m/%d/%Y}: {'; '.join(updated)}."
    if failed:
        note += f" REDCap update FAILED for {'; '.join(failed)}."
    return "\n".join(line for line in lines if line), note


def _number(value: str):
    """A single record ID goes in as a number (like the manual sheet); lists like '256, 312' stay text."""
    return int(value) if isinstance(value, str) and value.isdigit() else value


def _open_workbook(path: Path):
    """openpyxl, tolerating workbooks (e.g. from R's openxlsx) that list parts which don't exist."""
    try:
        return load_workbook(path)
    except KeyError:
        pass
    src = zipfile.ZipFile(path)
    names = set(src.namelist())
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as dst:
        for name in src.namelist():
            data = src.read(name)
            if name.endswith(".rels"):
                base = posixpath.dirname(posixpath.dirname(name))
                text = data.decode("utf-8")
                for target in re.findall(r'Target="([^"]+)"', text):
                    resolved = posixpath.normpath(posixpath.join(base, target)).lstrip("/")
                    if "://" not in target and resolved not in names:
                        text = re.sub(rf'<Relationship [^>]*Target="{re.escape(target)}"[^>]*/>', "", text)
                data = text.encode("utf-8")
            dst.writestr(name, data)
    buf.seek(0)
    return load_workbook(buf)


def write_tracking_tab(source: Path, output: Path, rows: list[list]) -> None:
    """Copy `source` to `output` (unless they're the same file) and refill its Tracking Duplicates tab.
    Every other sheet is left as it is. When writing over the source itself, a backup is made first."""
    if source.resolve() == output.resolve():
        backup = source.with_name(f"{source.stem}_backup_{date.today():%Y-%m-%d}{source.suffix}")
        shutil.copyfile(source, backup)
        print(f"Backup of {source.name}: {backup.name}")
    wb = _open_workbook(source)
    if SHEET in wb.sheetnames:
        ws = wb[SHEET]
        for row in ws.iter_rows(min_row=2, max_row=max(ws.max_row, 2)):
            for cell in row:
                cell.value = None
        for row in range(2, max(ws.max_row, 2) + 1):          # status colors from an earlier run
            ws.cell(row=row, column=STATUS_COL).fill = PatternFill()
        notes_header = ws.cell(row=1, column=NOTES_COL)
        for col in (AUTOMATION_COL, STATUS_COL):
            new_header = ws.cell(row=1, column=col)
            if new_header.value != HEADERS[col - 1]:
                # tabs from before these columns existed: add the header, styled like the notes header
                new_header.value = HEADERS[col - 1]
                if notes_header.has_style:
                    new_header.font, new_header.fill = copy(notes_header.font), copy(notes_header.fill)
                    new_header.border = copy(notes_header.border)
                    new_header.alignment = copy(notes_header.alignment)
    else:
        ws = wb.create_sheet(SHEET)
        for j, header in enumerate(HEADERS, start=1):
            cell = ws.cell(row=1, column=j, value=header)
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="D9E1F2")
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    wrap = Alignment(wrap_text=True, vertical="top")
    for i, values in enumerate(rows, start=2):
        for j, value in enumerate(values, start=1):
            cell = ws.cell(row=i, column=j, value=value if value != "" else None)
            if isinstance(value, date):
                cell.number_format = "mm/dd/yyyy"
            cell.alignment = wrap
            if j == STATUS_COL and value in STATUS_COLORS:
                cell.fill = PatternFill("solid", fgColor=STATUS_COLORS[value])
    ws.column_dimensions["B"].width = max(ws.column_dimensions["B"].width or 0, 40)
    ws.column_dimensions["M"].width = max(ws.column_dimensions["M"].width or 0, 60)
    ws.column_dimensions["N"].width = max(ws.column_dimensions["N"].width or 0, 90)
    ws.column_dimensions["O"].width = max(ws.column_dimensions["O"].width or 0, 16)
    ws.freeze_panes = "A2"
    wb.save(output)
