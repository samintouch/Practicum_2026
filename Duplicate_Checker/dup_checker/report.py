"""Write the review workbook."""
from __future__ import annotations

from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .checker import PairResult, RecordResult
from .tracking import HEADERS as TRACKING_HEADERS
from .tracking import build_rows as build_tracking_rows

FILLS = {
    "Duplicate": PatternFill("solid", fgColor="F8CBAD"),
    "Not a duplicate": PatternFill("solid", fgColor="C6EFCE"),
    "Needs review": PatternFill("solid", fgColor="FFEB9C"),
    "DUPLICATE": PatternFill("solid", fgColor="F8CBAD"),
    "NOT_DUPLICATE": PatternFill("solid", fgColor="C6EFCE"),
    "UNSURE": PatternFill("solid", fgColor="FFEB9C"),
    "MATCH": PatternFill("solid", fgColor="C6EFCE"),
    "DIFFERENT": PatternFill("solid", fgColor="F8CBAD"),
    "SUITE_DIFFERS": PatternFill("solid", fgColor="FFEB9C"),
    "UNREACHABLE": PatternFill("solid", fgColor="D9D9D9"),
    "updated": PatternFill("solid", fgColor="C6EFCE"),
    "partly updated": PatternFill("solid", fgColor="C6EFCE"),
    "need review": PatternFill("solid", fgColor="F4B183"),
    "Not Duplicate": PatternFill("solid", fgColor="E7E6E6"),
    "skipped": PatternFill("solid", fgColor="FFEB9C"),
    "failed": PatternFill("solid", fgColor="F8CBAD"),
}


def _yn(value: bool | None) -> str:
    return "" if value is None else ("Yes" if value else "No")


def _record_rows(results: list[RecordResult]) -> pd.DataFrame:
    rows = []
    for res in results:
        r, w = res.record, res.website
        rows.append({
            "duplicate_group_id": r.group_id,
            "record_id": r.record_id,
            "organization_name": r.name,
            "address": r.full_address,
            "phone": r.phone,
            "website": r.website,
            "recommendation": res.recommendation,
            "duplicate_with": ", ".join(res.duplicate_with),
            "primary_record": res.primary_id or "",
            "decided_by": res.decided_by,
            "confidence": round(res.confidence, 2) if res.confidence is not None else None,
            "reason": res.reason,
            "website_status": w.status if w else "",
            "website_detail": w.detail if w else "",
            "address_seen_on_website": (w.matched_text or "; ".join(w.found_addresses[:3])) if w else "",
            "suggested_address_change": res.suggested_address_change,
            "name_on_website": _yn(w.name_on_page) if w else "",
            "draft_note": res.draft_note,
            "completeness_score": r.completeness_score,
            "original_match_reason": r.duplicate_reason,
        })
    return pd.DataFrame(rows)


def _pair_rows(pairs: list[PairResult]) -> pd.DataFrame:
    return pd.DataFrame([{
        "duplicate_group_id": p.group_id,
        "record_a": p.a.record_id,
        "record_b": p.b.record_id,
        "name_a": p.a.name,
        "name_b": p.b.name,
        "address_a": p.a.full_address,
        "address_b": p.b.full_address,
        "name_match": p.name_level,
        "name_score": p.name_score,
        "shared_name_words": ", ".join(p.shared_words),
        "address_match": p.address_level,
        "address_detail": p.address_detail,
        "same_phone": _yn(p.phone_match),
        "same_zip": _yn(p.zip_match),
        "website_cross_check": p.cross_site,
        "verdict": p.verdict,
        "confidence": round(p.confidence, 2),
        "decided_by": p.decided_by,
        "reason": p.reason,
    } for p in pairs])


def _website_rows(results: list[RecordResult]) -> pd.DataFrame:
    rows = []
    for res in results:
        w = res.website
        if w is None:
            continue
        rows.append({
            "record_id": res.record.record_id,
            "organization_name": res.record.name,
            "redcap_address": res.record.full_address,
            "website": w.url,
            "final_url": w.final_url,
            "status": w.status,
            "detail": w.detail,
            "matched_text": w.matched_text,
            "addresses_found": "\n".join(w.found_addresses),
            "suggested_address": w.suggested_address,
            "found_by_ai": "Yes" if w.found_by_ai else "",
            "read_with": w.via,
            "name_on_website": _yn(w.name_on_page),
            "page_title": w.page_title,
            "pages_checked": "\n".join(w.pages_checked),
        })
    return pd.DataFrame(rows)


def _tracking_rows(results: list[RecordResult]) -> pd.DataFrame:
    return pd.DataFrame(build_tracking_rows(results), columns=TRACKING_HEADERS)


def _summary_rows(results: list[RecordResult], pairs: list[PairResult], settings: dict) -> pd.DataFrame:
    rec = pd.Series([("Duplicate" if r.recommendation.startswith("Duplicate") else r.recommendation)
                     for r in results]).value_counts()
    web = pd.Series([r.website.status for r in results if r.website]).value_counts()
    by = pd.Series([p.decided_by or "-" for p in pairs]).value_counts()
    rows = [("Run date", datetime.now().strftime("%Y-%m-%d %H:%M"))]
    rows += [(k, v) for k, v in settings.items()]
    rows += [("", ""), ("Groups checked", len({r.record.group_id for r in results})),
             ("Records checked", len(results)), ("Record pairs compared", len(pairs)), ("", "")]
    rows += [(f"Records: {k}", int(v)) for k, v in rec.items()]
    rows += [("", "")] + [(f"Pairs decided by: {k}", int(v)) for k, v in by.items()]
    rows += [("", "")] + [(f"Website: {k}", int(v)) for k, v in web.items()]
    rows += [("", ""), ("Note", "All results are suggestions. Verify before updating REDCap.")]
    return pd.DataFrame(rows, columns=["Item", "Value"])


def _redcap_rows(plans: list, applied: bool) -> pd.DataFrame:
    yes_no = {"1": "Yes", "0": "No", "": "(blank)"}
    rows = []
    for p in plans:
        row = {
            "record_id": p.record_id,
            "organization_name": p.organization_name,
            "role": "record kept (most complete)" if p.role == "kept" else f"duplicate of {p.primary_id}",
            "decided_by": p.decided_by,
            "set_duplicate_to": yes_no[p.set_duplicate],
            "set_validation_to": "Yes",
            "set_validation_date_to": p.validation_date,
            "comment_to_add": p.comment_to_add,
        }
        if applied:
            row.update({
                "status": p.status,
                "fields_changed": ", ".join(p.changes) if p.action == "UPDATE" else "",
                "note": p.note,
                "name_in_redcap": p.redcap_name,
                "old_duplicate": yes_no.get(p.old_duplicate, p.old_duplicate),
                "old_validation": yes_no.get(p.old_validated, p.old_validated),
                "old_validation_date": p.old_validation_date,
                "old_comments": p.old_comments,
            })
        else:
            row["status"] = "preview - run with --apply to update REDCap"
        rows.append(row)
    return pd.DataFrame(rows)


def write_redcap_log(path: Path, plans: list) -> None:
    """What --apply changed in REDCap, with each record's old values (for undo)."""
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        _redcap_rows(plans, applied=True).to_excel(writer, sheet_name="REDCap Updates", index=False)
    _style(path)


def write_report(path: Path, results: list[RecordResult], pairs: list[PairResult], settings: dict,
                 redcap_plans: list | None = None, redcap_applied: bool = False) -> None:
    sheets = {
        "Summary": _summary_rows(results, pairs, settings),
        "Record Recommendations": _record_rows(results),
        "Pair Comparisons": _pair_rows(pairs),
        "Website Checks": _website_rows(results),
        "Tracking Duplicates (Draft)": _tracking_rows(results),
    }
    if redcap_plans is not None:
        sheets["REDCap Updates"] = _redcap_rows(redcap_plans, redcap_applied)
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        for name, df in sheets.items():
            df.to_excel(writer, sheet_name=name, index=False)
    _style(path)


def _style(path: Path) -> None:
    wb = load_workbook(path)
    wide = {"reason", "draft_note", "website_detail", "address_seen_on_website", "addresses_found",
            "notes (add to redcap)", "pages_checked", "suggested_address_change", "address_detail",
            "comment_to_add", "old_comments", "note"}
    color_cols = {"recommendation", "verdict", "status", "website_status"}
    for ws in wb.worksheets:
        ws.freeze_panes = "A2" if ws.title != "Summary" else None
        if ws.max_row > 1 and ws.title != "Summary":
            ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="D9E1F2")
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        for col in ws.iter_cols(min_row=1):
            header = str(col[0].value or "")
            longest = max((len(str(c.value)) for c in col if c.value is not None), default=10)
            width = 60 if header in wide or header.startswith("notes") else min(max(longest, len(header) * 0.6, 10) + 2, 45)
            ws.column_dimensions[col[0].column_letter].width = width
            for c in col[1:]:
                c.alignment = Alignment(wrap_text=width >= 45, vertical="top")
                if header in color_cols and c.value:
                    key = "Duplicate" if str(c.value).startswith("Duplicate") else str(c.value)
                    if key in FILLS:
                        c.fill = FILLS[key]
    wb.save(path)
